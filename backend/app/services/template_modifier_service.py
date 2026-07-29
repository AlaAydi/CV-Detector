import re
import shutil
from pathlib import Path

import fitz
from docx import Document
from docx.text.paragraph import Paragraph

from app.services.optimizer_service import ResumeSection, _match_section_key, optimize_sections, sections_have_changes


def _detect_bullet_prefix(text: str) -> str:
    stripped = text.lstrip()
    if not stripped:
        return ""
    if stripped[0] in "-•·*▪":
        return stripped[0] + " "
    if stripped.startswith(("– ", "— ")):
        return stripped[:2]
    return ""


def _title_variants(title: str) -> list[str]:
    return list(dict.fromkeys([title, title.title(), title.upper(), title.lower()]))


def _token_is_highlighted(text: str, highlight_tokens: frozenset[str]) -> bool:
    lower = text.lower()
    return any(token in lower for token in highlight_tokens)


def _clear_paragraph_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        element = run._element
        element.getparent().remove(element)


def _copy_run_font(run, reference_run) -> None:
    if reference_run is None:
        return
    run.font.name = reference_run.font.name
    run.font.size = reference_run.font.size
    if reference_run.font.color.rgb is not None:
        run.font.color.rgb = reference_run.font.color.rgb


def _write_inline_with_highlights(
    paragraph: Paragraph,
    line: str,
    highlight_tokens: frozenset[str],
    reference_run=None,
) -> None:
    separator = ", " if "," in line else "; "
    parts = [part.strip() for part in re.split(r"[,;]", line) if part.strip()]
    _clear_paragraph_runs(paragraph)

    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        _copy_run_font(run, reference_run)
        if _token_is_highlighted(part, highlight_tokens):
            run.bold = True
        elif reference_run is not None:
            run.bold = reference_run.bold
        if index < len(parts) - 1:
            sep_run = paragraph.add_run(separator)
            _copy_run_font(sep_run, reference_run)


def _write_line_with_highlights(
    paragraph: Paragraph,
    line: str,
    highlight_tokens: frozenset[str],
    bullet_prefix: str,
    reference_run=None,
) -> None:
    display = line
    if bullet_prefix and not line.lstrip().startswith(("-", "•", "·", "*", "▪", "–", "—")):
        display = bullet_prefix + line

    _clear_paragraph_runs(paragraph)
    run = paragraph.add_run(display)
    _copy_run_font(run, reference_run)
    if _token_is_highlighted(line, highlight_tokens):
        run.bold = True
    elif reference_run is not None:
        run.bold = reference_run.bold


def _write_lines_to_paragraphs(
    paragraphs,
    indices: list[int],
    section: ResumeSection,
    bullet_prefix: str,
) -> None:
    lines = section.lines
    highlight = section.highlight_tokens
    inline = len(indices) == 1 and len(lines) == 1 and ("," in lines[0] or ";" in lines[0])
    reference_run = paragraphs[indices[0]].runs[0] if paragraphs[indices[0]].runs else None

    if inline:
        _write_inline_with_highlights(paragraphs[indices[0]], lines[0], highlight, reference_run)
        return

    for index, para_idx in enumerate(indices):
        if index < len(lines):
            _write_line_with_highlights(
                paragraphs[para_idx],
                lines[index],
                highlight,
                bullet_prefix,
                reference_run,
            )
        else:
            paragraphs[para_idx].text = ""

    if len(lines) > len(indices):
        anchor = paragraphs[indices[-1]]
        for line in lines[len(indices) :]:
            text = bullet_prefix + line if bullet_prefix else line
            new_para = anchor.insert_paragraph_after(text)
            if _token_is_highlighted(line, highlight):
                for run in new_para.runs:
                    run.bold = True
            anchor = new_para


def _apply_sections_to_docx(src: Path, dst: Path, sections: list[ResumeSection]) -> bool:
    doc = Document(src)
    paragraphs = doc.paragraphs
    sections_by_key = {section.key: section for section in sections if section.key != "header"}
    modified = False
    index = 0

    while index < len(paragraphs):
        text = paragraphs[index].text.strip()
        if not text:
            index += 1
            continue

        section_key = _match_section_key(text)
        if not section_key:
            index += 1
            continue

        optimized = sections_by_key.get(section_key)
        if not optimized or not optimized.lines:
            index += 1
            continue

        body_start = index + 1
        body_end = body_start
        while body_end < len(paragraphs):
            candidate = paragraphs[body_end].text.strip()
            if candidate and _match_section_key(candidate):
                break
            body_end += 1

        body_indices = [i for i in range(body_start, body_end) if paragraphs[i].text.strip()]
        if not body_indices:
            index = body_end
            continue

        bullet_prefix = _detect_bullet_prefix(paragraphs[body_indices[0]].text)
        _write_lines_to_paragraphs(paragraphs, body_indices, optimized, bullet_prefix)
        modified = True
        index = body_end

    doc.save(dst)
    return modified


def _find_section_end_y(page: fitz.Page, y_start: float) -> float:
    y_end = page.rect.height - 40
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            line_text = "".join(span["text"] for span in line["spans"]).strip()
            line_y = line["bbox"][1]
            if line_y <= y_start:
                continue
            if _match_section_key(line_text):
                y_end = min(y_end, line_y - 2)
    return y_end


def _find_title_rect(page: fitz.Page, title: str) -> fitz.Rect | None:
    for variant in _title_variants(title):
        rects = page.search_for(variant)
        if rects:
            return rects[0]
    normalized_title = re.sub(r"[^a-z0-9àâäéèêëïîôùûüç\s]", "", title.lower()).strip()
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            line_text = "".join(span["text"] for span in line["spans"]).strip()
            normalized_line = re.sub(r"[^a-z0-9àâäéèêëïîôùûüç\s]", "", line_text.lower()).strip()
            if normalized_line == normalized_title or normalized_line.startswith(normalized_title):
                return fitz.Rect(line["bbox"])
    return None


def _get_body_line_entries(page: fitz.Page, y_start: float, y_end: float) -> list[tuple[fitz.Rect, str, list]]:
    entries: list[tuple[fitz.Rect, str, list]] = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            bbox = line["bbox"]
            if bbox[1] < y_start or bbox[1] >= y_end:
                continue
            text = "".join(span["text"] for span in line["spans"]).strip()
            if not text or _match_section_key(text):
                continue
            entries.append((fitz.Rect(bbox), text, line["spans"]))
    return sorted(entries, key=lambda item: item[0].y0)


def _span_style(spans: list) -> dict:
    for span in spans:
        if span.get("text", "").strip():
            flags = span.get("flags", 0)
            return {
                "size": float(span.get("size", 11)),
                "font": span.get("font", "helv"),
                "color": span.get("color", 0),
                "bold": bool(flags & 2**4),
            }
    return {"size": 11.0, "font": "helv", "color": 0, "bold": False}


def _pdf_fontname(style: dict, bold: bool) -> str:
    if bold:
        return "hebo"
    font_lower = style.get("font", "").lower()
    if "bold" in font_lower:
        return "hebo"
    if "italic" in font_lower or "oblique" in font_lower:
        return "heit"
    return "helv"


def _color_from_int(color: int) -> tuple:
    r = ((color >> 16) & 255) / 255
    g = ((color >> 8) & 255) / 255
    b = (color & 255) / 255
    return (r, g, b)


def _write_line_on_pdf(
    page: fitz.Page,
    rect: fitz.Rect,
    line: str,
    highlight_tokens: frozenset[str],
    style: dict,
) -> None:

    # Effacer uniquement l'ancien contenu
    area = fitz.Rect(
        rect.x0 - 1,
        rect.y0 - 1,
        rect.x1 + 1,
        rect.y1 + 2
    )

    page.add_redact_annot(area, fill=(1, 1, 1))
    page.apply_redactions()


    fontsize = style.get("size", 10)
    color = _color_from_int(style.get("color", 0))


    # largeur maximale de la zone originale
    max_width = rect.width


    skills = []

    if "," in line or ";" in line:
        skills = [
            x.strip()
            for x in re.split(r"[,;]", line)
            if x.strip()
        ]
    else:
        skills = [line]


    x = rect.x0
    y = rect.y1 - 2


    for index, skill in enumerate(skills):

        separator = ", " if index < len(skills)-1 else ""

        text = skill + separator


        bold = _token_is_highlighted(
            skill,
            highlight_tokens
        )


        font = _pdf_fontname(
            style,
            bold
        )


        width = fitz.get_text_length(
            text,
            fontname=font,
            fontsize=fontsize
        )


        # si dépassement de la ligne originale
        if x + width > rect.x0 + max_width:

            # On arrête pour ne pas casser le template
            break


        page.insert_text(
            (x, y),
            text,
            fontname=font,
            fontsize=fontsize,
            color=color
        )


        x += width

def _apply_section_to_pdf_page(
    page: fitz.Page,
    section: ResumeSection
) -> bool:

    if not section.title or not section.lines:
        return False


    header_rect = _find_title_rect(
        page,
        section.title
    )


    if not header_rect:
        return False


    y_start = header_rect.y1 + 2

    y_end = _find_section_end_y(
        page,
        y_start
    )


    body_entries = _get_body_line_entries(
        page,
        y_start,
        y_end
    )


    if not body_entries:
        return False


    modified = False


    # On garde uniquement l'espace existant
    available = len(body_entries)


    new_lines = section.lines[:available]


    for index, optimized_line in enumerate(new_lines):

        rect, original, spans = body_entries[index]


        if optimized_line.strip() == original.strip():
            continue


        style = _span_style(spans)


        _write_line_on_pdf(
            page,
            rect,
            optimized_line,
            section.highlight_tokens,
            style
        )


        modified = True


    return modified


def _apply_sections_to_pdf(src: Path, dst: Path, sections: list[ResumeSection]) -> bool:
    doc = fitz.open(src)
    modified = False

    try:
        for section in sections:
            if section.key == "header" or not section.title or not section.lines:
                continue
            for page in doc:
                if _apply_section_to_pdf_page(page, section):
                    modified = True
                    break

        doc.save(dst)
    finally:
        doc.close()

    return modified


def modify_resume_preserving_template(
    original_path: Path,
    output_path: Path,
    resume_text: str,
    job_text: str,
) -> bool:


    sections = optimize_sections(
        resume_text,
        job_text
    )


    # uniquement compétences
    sections = [
        s for s in sections
        if s.key == "skills"
    ]


    if not sections:
        shutil.copy2(
            original_path,
            output_path
        )
        return False


    has_changes = sections_have_changes(
        sections
    )


    if original_path.suffix.lower() == ".pdf":

        modified = _apply_sections_to_pdf(
            original_path,
            output_path,
            sections
        )

    elif original_path.suffix.lower() == ".docx":

        modified = _apply_sections_to_docx(
            original_path,
            output_path,
            sections
        )

    else:

        shutil.copy2(
            original_path,
            output_path
        )

        return False



    if not modified:

        shutil.copy2(
            original_path,
            output_path
        )


    return modified and has_changes